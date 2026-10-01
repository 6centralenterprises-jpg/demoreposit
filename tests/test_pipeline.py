import copy
import json
import shutil
from argparse import Namespace
from pathlib import Path

import pytest
from openpyxl import load_workbook

import qualify
from partner_qualifier.normalize import normalize, read_rows
from partner_qualifier.rubric import score

FIXTURES = Path(__file__).parent / "fixtures"
MARKET = "Chicago, IL"


@pytest.fixture
def leads():
    return {l["lead_id"]: l for l in normalize(read_rows(FIXTURES / "sample_igleads.csv"), MARKET)}


def evidence(lead_id):
    return json.loads((FIXTURES / "evidence" / f"{lead_id}.json").read_text())


def test_normalize_cleans_and_dedupes(leads):
    assert list(leads) == ["L0001", "L0002", "L0003", "L0004", "L0005", "L0006"]  # blank row dropped
    assert leads["L0002"]["screen"] == "skip_duplicate"
    assert leads["L0002"]["duplicate_of"] == "L0001"  # same email, different case
    assert leads["L0001"]["email_type"] == "business_domain"
    assert leads["L0003"]["email_type"] == "free_mail"
    assert leads["L0001"]["phone"] == "(773) 555-0101"
    assert leads["L0003"]["phone"] == "(312) 555-0123"  # pulled from bio
    assert leads["L0001"]["followers"] == 1200
    assert leads["L0001"]["website"] == "https://sparkleexample.test"  # pulled from bio
    assert leads["L0005"]["website_kind"] == "link_page"
    assert leads["L0001"]["niche_guess"] == "house cleaning"
    assert leads["L0003"]["niche_guess"] == "car detailing"
    assert leads["L0004"]["niche_guess"] == "pest control"
    assert leads["L0001"]["location_hint"] == "strong"
    assert leads["L0005"]["location_hint"] == "none"
    assert leads["L0001"]["priority"] > leads["L0006"]["priority"]


def test_fully_sourced_business_is_message(leads):
    result = score(leads["L0001"], evidence("L0001"), MARKET)
    assert result["verdict"] == "MESSAGE"
    assert result["score"] == 10
    assert result["contact"]["email_basis"] == "publicly listed"
    assert result["opener"]["text"]


def test_unsourced_claims_never_count(leads):
    result = score(leads["L0003"], evidence("L0003"), MARKET)
    insured = next(b for b in result["breakdown"] if b["key"] == "insured")
    assert insured["status"] == "unknown" and insured["points"] == 0  # "insured" had no source
    assert result["score"] == 4  # rating 2 + presence 1 + serves target 1
    assert result["verdict"] == "VERIFY"
    assert "insurance" in result["to_verify"]
    assert result["reason"].startswith("4/10 confirmed; could reach 10")
    assert result["contact"]["email_free_mail"] is True


def test_pest_control_without_verified_license_is_do_not_contact(leads):
    result = score(leads["L0004"], evidence("L0004"), MARKET)
    assert result["verdict"] == "DO NOT CONTACT"
    ev = evidence("L0004")
    ev["license"] = {"required": True, "number": "SPC-0000", "status": "active",
                     "source": "https://license.example/0000"}
    assert score(leads["L0004"], ev, MARKET)["verdict"] == "MESSAGE"


def test_gates_and_flags(leads):
    assert score(leads["L0005"], evidence("L0005"), MARKET)["verdict"] == "SKIP"
    assert score(leads["L0006"], evidence("L0006"), MARKET)["reason"] == "Not a local service provider"
    assert score(leads["L0002"], None, MARKET)["verdict"] == "SKIP"
    assert score(leads["L0001"], None, MARKET)["verdict"] == "NOT RESEARCHED"

    flagged = copy.deepcopy(evidence("L0001"))
    flagged["red_flags"] = [{"flag": "Open lawsuit over property damage", "severity": "high",
                             "source": "https://court.example/case"}]
    assert score(leads["L0001"], flagged, MARKET)["verdict"] == "HOLD"
    flagged["red_flags"][0]["source"] = None  # unsourced concern: noted, not a HOLD
    result = score(leads["L0001"], flagged, MARKET)
    assert result["verdict"] == "MESSAGE"
    assert result["unsourced_concerns"] == ["Open lawsuit over property damage"]


def test_low_ceiling_is_skip(leads):
    ev = evidence("L0003")
    ev["reviews"] = [{"platform": "Google", "rating": 3.1, "count": 60, "source": "https://g.example/x"}]
    ev["size"] = {"value": "50+", "source": "https://g.example/x"}
    ev["insured"] = {"value": False, "source": "https://g.example/x"}
    ev["activity"] = {"last_active": "2024-01", "source": "https://g.example/x"}
    result = score(leads["L0003"], ev, MARKET)
    assert result["max_possible"] < 7
    assert result["verdict"] == "SKIP"


def test_end_to_end(tmp_path, monkeypatch):
    monkeypatch.setattr(qualify, "RUNS", tmp_path)
    qualify.cmd_prepare(Namespace(file=str(FIXTURES / "sample_igleads.csv"), market=MARKET,
                                  run="test", batch_size=2, limit=None))
    run = tmp_path / "test"
    batches = sorted((run / "batches").glob("batch_*.json"))
    queued = [l["lead_id"] for b in batches for l in json.loads(b.read_text())["leads"]]
    assert queued[0] == "L0001" and "L0002" not in queued and len(queued) == 5
    for f in (FIXTURES / "evidence").glob("*.json"):
        shutil.copy(f, run / "evidence" / f.name)
    qualify.cmd_score(Namespace(run="test", out_dir=str(tmp_path / "out")))

    verdicts = {r["lead_id"]: r["verdict"] for r in json.loads((run / "results.json").read_text())}
    assert verdicts == {"L0001": "MESSAGE", "L0002": "SKIP", "L0003": "VERIFY",
                        "L0004": "DO NOT CONTACT", "L0005": "SKIP", "L0006": "SKIP"}
    wb = load_workbook(run / "partner_qualification.xlsx")
    assert wb.sheetnames == ["Summary", "Who to Message", "All Leads", "Evidence", "Partner Sheet"]
    assert wb["Who to Message"]["B2"].value == "Sparkle Example Cleaning LLC"
    assert wb["Partner Sheet"]["A1"].value == "Business"
    assert (tmp_path / "out" / "test_partner_qualification.xlsx").exists()


def test_bare_domain_in_bio_becomes_website():
    rows = [{"username": "shiny", "bio": "Book at shinyhomes.example.com or hi@shinyhomes.example.com"}]
    lead = normalize(rows, MARKET)[0]
    assert lead["website"] == "https://shinyhomes.example.com"
    assert lead["email"] == "hi@shinyhomes.example.com"

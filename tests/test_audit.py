import json
import shutil
from argparse import Namespace
from pathlib import Path

from openpyxl import load_workbook

import qualify
from partner_qualifier.audit import (OPT_OUT, audit_partner, check_email, local_leaders, local_pack_hits,
                                     missing_keywords, visibility)

FIXTURES = Path(__file__).parent / "fixtures"
SENDER = {"name": "Terell John", "title": "Founder", "company": "6 Central Enterprises",
          "phone": "[YOUR PHONE]", "mailing_address": "[YOUR MAILING ADDRESS]"}


def load(path):
    return json.loads((FIXTURES / path).read_text())


def test_market_leaders_and_map_pack_match():
    market = load("market/house-cleaning.json")
    leaders = local_leaders(market)
    assert leaders[0]["name"] == "Big Example Maids" and leaders[0]["appearances"] == 2
    hits = local_pack_hits("sparkleexample.test", "Sparkle Example Cleaning LLC", market)
    assert [h["keyword"] for h in hits] == ["house cleaning chicago"]
    assert local_pack_hits("", "Sparkle Example Cleaning LLC", market)  # name match without a domain


def test_missing_keywords_skip_page_one_and_two():
    missing = missing_keywords(load("audits/L0001.json"), load("market/house-cleaning.json"))
    assert [k["keyword"] for k in missing] == ["move out cleaning chicago", "maid service chicago"]


def test_visibility_no_website_means_high_need():
    audit = {"seo": {"website_type": "none"}, "google_reputation": {}}
    market = load("market/house-cleaning.json")
    score, parts = visibility(audit, {"reviews": [{"count": 40, "source": "https://x.example"}]}, market, [])
    assert score == 0  # no site, no rankings, no map pack, no visitors, 40 reviews vs leader median 600
    assert all(p["known"] for p in parts)


def test_good_email_passes_and_gets_signature():
    result = check_email(load("audits/L0001.json")["email"], SENDER)
    assert result["status"] == "Draft: needs Terell's approval", result["failed"]
    assert result["full_text"].endswith(OPT_OUT)
    assert "Founder, 6 Central Enterprises" in result["full_text"]
    assert result["words"] < 120


def test_bad_email_needs_rewrite():
    email = {"subject": "URGENT: exclusive leads for you, act now before spots are filling up",
             "body": "Hi! We guarantee 20 jobs a week and steady work. Want more income? Interested?",
             "facts_used": [{"fact": "no source"}]}
    result = check_email(email, SENDER)
    assert result["status"] == "Needs rewrite"
    failed = " ".join(result["failed"])
    for problem in ("Says who we are", "Price agreed up front", "Exactly one question",
                    "Uses a true, sourced fact", "No risky claims", "Short subject line"):
        assert problem in failed
    for claim in ("guarantee", "exclusivity", "income promise", "job-volume promise", "fake urgency"):
        assert claim in failed


def test_audit_partner_scores_fit():
    lead = {"lead_id": "L0001", "website": "https://sparkleexample.test"}
    result = {"score": 10, "verdict": "MESSAGE", "business_name": "Sparkle Example Cleaning LLC",
              "niche": "house cleaning", "contact": {"best": "Email hello@sparkleexample.test"},
              "website": "https://sparkleexample.test"}
    row = audit_partner(lead, load("evidence/L0001.json"), result, load("audits/L0001.json"),
                        load("market/house-cleaning.json"), SENDER)
    # Visibility: site 2 + page-1 keywords 1 + map pack 1 + visitors 1 + reviews 1 (215 vs median 600) = 6
    assert row["visibility"] == 6
    assert row["need"] == 5  # 10 - 6 + 1 capacity signal
    assert row["quadrant"] == "Good partner"
    assert row["google"]["via"] == "Google map pack (live)" and row["google"]["reviews"] == 215
    assert len(row["doing_wrong"]) == 1 and row["unsourced_points"] == 1  # unsourced point dropped
    assert row["competes_on_google"] is True  # in the map pack for 1 of 2 searches checked


def test_audit_end_to_end(tmp_path, monkeypatch):
    monkeypatch.setattr(qualify, "RUNS", tmp_path)
    monkeypatch.setattr(qualify, "CONFIG", tmp_path / "no-config")
    qualify.cmd_prepare(Namespace(file=str(FIXTURES / "sample_igleads.csv"), market="Chicago, IL",
                                  run="t", batch_size=5, limit=None))
    run = tmp_path / "t"
    for f in (FIXTURES / "evidence").glob("*.json"):
        shutil.copy(f, run / "evidence" / f.name)
    qualify.cmd_score(Namespace(run="t", out_dir=None))
    qualify.cmd_audit_plan(Namespace(run="t", top=20, include_verify=True, batch_size=3))
    plan = json.loads((run / "audit_plan.json").read_text())
    assert plan["partners"] == ["L0001", "L0003"]  # MESSAGE first, then VERIFY that could reach 8+
    assert plan["missing_markets"] == ["house cleaning", "car detailing"]
    batch = json.loads((run / "audit_batches" / "batch_01.json").read_text())
    assert batch["partners"][0]["domain_to_audit"] == "sparkleexample.test"
    assert batch["partners"][0]["market_file"].endswith("market/house-cleaning.json")

    shutil.copy(FIXTURES / "market" / "house-cleaning.json", run / "market" / "house-cleaning.json")
    shutil.copy(FIXTURES / "audits" / "L0001.json", run / "audits" / "L0001.json")
    qualify.cmd_audit_report(Namespace(run="t", out_dir=str(tmp_path / "out")))
    wb = load_workbook(run / "partner_audits.xlsx")
    assert wb.sheetnames == ["Best Partners", "SEO & Keywords", "Market Leaders", "Demand Keywords",
                             "Right & Wrong", "Email Drafts"]
    assert wb["Best Partners"]["B2"].value == "Sparkle Example Cleaning LLC"
    assert wb["Email Drafts"]["F2"].value == "Draft: needs Terell's approval"
    book = (run / "partner_audit_book.md").read_text()
    for heading in ("### Google reputation", "### SEO", "### Keywords", "### Competition",
                    "### What they're doing right", "### What they're doing wrong", "### Custom email"):
        assert heading in book
    assert (tmp_path / "out" / "t_partner_audit_book.md").exists()


def test_partner_without_matching_market_is_not_compared_to_another_niche():
    from partner_qualifier.audit import market_for
    markets = {"house cleaning": load("market/house-cleaning.json")}
    assert market_for("house cleaning", markets) is markets["house cleaning"]
    assert market_for("car detailing", markets) is None

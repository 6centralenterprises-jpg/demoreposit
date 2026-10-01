"""Build the ranked spreadsheet and the paste-ready partner sheet."""
import csv
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .rubric import CRITERIA, MESSAGE_THRESHOLD

VERDICT_ORDER = ["MESSAGE", "VERIFY", "HOLD", "DO NOT CONTACT", "SKIP", "NOT RESEARCHED"]
VERDICT_FILL = {
    "MESSAGE": "C6EFCE", "VERIFY": "FFF2CC", "HOLD": "F8CBAD",
    "DO NOT CONTACT": "F4B6B6", "SKIP": "E7E6E6", "NOT RESEARCHED": "DDEBF7",
}
# Maps a verdict onto the partner sheet's status list.
SHEET_STATUS = {
    "MESSAGE": "New", "VERIFY": "New", "HOLD": "New", "NOT RESEARCHED": "New",
    "SKIP": "Not a Fit", "DO NOT CONTACT": "Do Not Contact",
}


def rank(results):
    def key(r):
        return (VERDICT_ORDER.index(r["verdict"]), -(r["score"] or 0), -(r["max_possible"] or 0),
                -r["review_total"])
    return sorted(results, key=key)


def _sheet(wb, title, headers, rows, widths=None, verdict_col=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F3864")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
    for i, header in enumerate(headers, start=1):
        width = (widths or {}).get(header, min(max(len(header) + 2, 12), 40))
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if verdict_col is not None:
            fill = VERDICT_FILL.get(row[verdict_col].value)
            if fill:
                row[verdict_col].fill = PatternFill("solid", fgColor=fill)
    ws.freeze_panes = "B2" if verdict_col is not None else "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    return ws


def _criterion_cell(breakdown, key):
    for b in breakdown:
        if b["key"] == key:
            mark = "?" if b["status"] == "unknown" else b["points"]
            return f"{mark}/{b['max']}"
    return ""


def build_workbook(path, leads, results, run):
    by_id = {lead["lead_id"]: lead for lead in leads}
    ranked = rank(results)
    wb = Workbook()
    wb.remove(wb.active)

    counts = Counter(r["verdict"] for r in results)
    summary = [
        ["Run", run["run_id"]],
        ["Market", run["market"]],
        ["Source file", run["source_file"]],
        ["Prepared", run["created_at"]],
        ["Leads in file", len(leads)],
        *[[f"Verdict: {v}", counts.get(v, 0)] for v in VERDICT_ORDER],
        [],
        ["How to read this", ""],
        ["MESSAGE", f"Scored {MESSAGE_THRESHOLD}+ out of 10 with sources. Contact these first, in rank order."],
        ["VERIFY", "Not enough confirmed yet, but could reach 7+. 'To verify' lists what to check."],
        ["HOLD", "Found a sourced red flag. Terell decides before anyone reaches out."],
        ["DO NOT CONTACT", "Needs a state license that could not be verified."],
        ["SKIP", "Duplicate, not a real business, out of area, or can't reach 7 even if unknowns check out."],
        ["?/n in a score column", "Not confirmed — earns 0 points (we never guess)."],
        [],
        ["Rubric (10 pts)", ""],
        *[[f"{area}: {label}", f"{pts} pt{'s' if pts > 1 else ''}"] for _, area, pts, label, _ in CRITERIA],
        ["Gate", "License-required niches (pest control, plumbing, roofing in IL) need a verified active license"],
    ]
    ws = wb.create_sheet("Summary")
    for row in summary:
        ws.append(row)
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 90
    for row in ws.iter_rows():
        if row[0].value in VERDICT_FILL:
            row[0].fill = PatternFill("solid", fgColor=VERDICT_FILL[row[0].value])
        if row[0].value in ("How to read this", "Rubric (10 pts)"):
            row[0].font = Font(bold=True)

    message_headers = ["Rank", "Business", "Score", "Niche", "Best contact", "Email", "Email basis",
                       "Phone", "Instagram", "Website", "Why they qualify", "Conversation opener",
                       "Opener source", "Notes"]
    message_rows = []
    for i, r in enumerate([r for r in ranked if r["verdict"] == "MESSAGE"], start=1):
        lead = by_id[r["lead_id"]]
        message_rows.append([
            i, r["business_name"], f"{r['score']}/10", r["niche"], r["contact"]["best"],
            r["contact"]["email"], r["contact"]["email_basis"], r["contact"]["phone"],
            lead["profile_url"], r["website"], r["why"], (r["opener"] or {}).get("text", ""),
            (r["opener"] or {}).get("source", ""), r["notes"],
        ])
    _sheet(wb, "Who to Message", message_headers, message_rows,
           widths={"Business": 28, "Best contact": 34, "Why they qualify": 60,
                   "Conversation opener": 50, "Notes": 40, "Email basis": 22})

    keys = [c[0] for c in CRITERIA]
    all_headers = ["Lead ID", "Verdict", "Score", "Max possible", "Business", "Instagram", "Niche",
                   "Reason", "To verify", *[c[3] for c in CRITERIA], "Red flags",
                   "Unsourced concerns", "License", "Best contact", "Website", "Followers",
                   "Location hint", "Researched"]
    all_rows = []
    for r in ranked:
        lead = by_id[r["lead_id"]]
        lic = r["license"] or {}
        all_rows.append([
            r["lead_id"], r["verdict"], r["score"], r["max_possible"], r["business_name"],
            lead["profile_url"], r["niche"], r["reason"], ", ".join(r["to_verify"]),
            *[_criterion_cell(r["breakdown"], k) for k in keys],
            "; ".join(f"{f.get('flag')} ({f.get('source')})" for f in r["red_flags"]),
            "; ".join(r["unsourced_concerns"]),
            (f"{lic.get('number') or ''} {lic.get('status') or ''}".strip() or "not verified")
            if r["license_required"] else "n/a",
            r["contact"]["best"], r["website"], lead.get("followers"), lead.get("location_hint"),
            r["researched_at"],
        ])
    _sheet(wb, "All Leads", all_headers, all_rows, verdict_col=1,
           widths={"Business": 28, "Reason": 50, "To verify": 40, "Red flags": 40,
                   "Best contact": 34, "Instagram": 30, "Website": 30})

    evidence_rows = []
    for r in ranked:
        for b in r["breakdown"]:
            points = "not confirmed" if b["status"] == "unknown" else f"{b['points']}/{b['max']}"
            evidence_rows.append([r["lead_id"], r["business_name"], b["area"], b["label"], points,
                                  b["finding"], b["source"]])
        if r["license_required"]:
            lic = r["license"] or {}
            evidence_rows.append([r["lead_id"], r["business_name"], "Gate", r["license_required"],
                                  "verified" if r["verdict"] != "DO NOT CONTACT" else "NOT verified",
                                  f"{lic.get('number') or ''} {lic.get('status') or ''}".strip(),
                                  lic.get("source", "")])
    _sheet(wb, "Evidence", ["Lead ID", "Business", "Area", "Criterion", "Points", "Finding", "Source"],
           evidence_rows, widths={"Business": 28, "Criterion": 36, "Finding": 60, "Source": 60})

    sheet_rows = partner_sheet_rows(leads, ranked, run)
    _sheet(wb, "Partner Sheet", PARTNER_SHEET_HEADERS, sheet_rows,
           widths={"Business": 28, "Notes": 70, "Email": 32})
    wb.save(path)


PARTNER_SHEET_HEADERS = ["Business", "Niche", "City", "Email", "Score", "Status",
                         "Last contact date", "Notes"]


def partner_sheet_rows(leads, ranked, run):
    rows = []
    for r in ranked:
        note = f"{r['verdict']}: {r['reason']}"
        if r["contact"]["best"] and not r["contact"]["email"]:
            note += f" | Contact: {r['contact']['best']}"
        rows.append([r["business_name"], r["niche"], run["market"], r["contact"]["email"],
                     r["score"] if r["score"] is not None else "", SHEET_STATUS[r["verdict"]], "",
                     note])
    return rows


def write_partner_csv(path, leads, results, run):
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.writer(f)
        writer.writerow(PARTNER_SHEET_HEADERS)
        writer.writerows(partner_sheet_rows(leads, rank(results), run))

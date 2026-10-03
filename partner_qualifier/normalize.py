"""Read an IGLeads (or similar) lead export and turn it into clean lead records.

Handles CSV, TSV, XLSX, JSON and JSONL. Column names vary between exports, so
headers are matched against a list of known aliases; anything unrecognized is
kept under "extra" so no data is lost.
"""
import csv
import json
import re
from pathlib import Path
from urllib.parse import urlparse

FIELD_ALIASES = {
    "handle": ["username", "user_name", "handle", "instagram_username", "ig_username",
               "instagram_handle", "account", "user"],
    "name": ["full_name", "fullname", "name", "display_name", "business_name", "company",
             "company_name"],
    "email": ["email", "public_email", "business_email", "contact_email", "emails",
              "email_address"],
    "phone": ["phone", "phone_number", "public_phone_number", "contact_phone_number",
              "business_phone", "phone_numbers", "mobile", "public_phone"],
    "bio": ["bio", "biography", "description", "about"],
    "website": ["external_url", "website", "url_in_bio", "bio_link", "link", "external_link",
                "site", "website_url"],
    "followers": ["followers", "follower_count", "followers_count"],
    "following": ["following", "following_count"],
    "posts": ["posts", "media_count", "posts_count", "post_count"],
    "verified": ["verified", "is_verified"],
    "profile_url": ["profile_url", "instagram_url", "profile_link", "profile", "url"],
    "category": ["category", "business_category", "category_name", "business_category_name"],
    "location": ["city", "address", "location", "business_address", "city_name", "state",
                 "zip", "zipcode", "zip_code", "country"],
    "is_business": ["is_business", "is_business_account", "business_account"],
    "source": ["source", "keyword", "search_term", "hashtag"],
}

FREE_MAIL_DOMAINS = {
    "gmail.com", "googlemail.com", "yahoo.com", "ymail.com", "rocketmail.com", "hotmail.com",
    "outlook.com", "live.com", "msn.com", "aol.com", "icloud.com", "me.com", "mac.com",
    "comcast.net", "att.net", "sbcglobal.net", "verizon.net", "protonmail.com", "proton.me",
    "gmx.com", "mail.com", "zoho.com", "yandex.com",
}

LINK_PAGE_DOMAINS = {
    "linktr.ee", "linkin.bio", "beacons.ai", "lnk.bio", "msha.ke", "taplink.cc", "bio.link",
    "campsite.bio", "solo.to", "hoo.be", "stan.store", "linkbio.co", "allmylinks.com",
}
BOOKING_DOMAINS = {
    "booksy.com", "vagaro.com", "styleseat.com", "calendly.com", "squareup.com", "setmore.com",
    "acuityscheduling.com", "housecallpro.com", "jobber.com", "getjobber.com", "wa.me",
}
SOCIAL_DOMAINS = {
    "instagram.com", "facebook.com", "fb.com", "tiktok.com", "youtube.com", "x.com",
    "twitter.com", "linkedin.com", "yelp.com", "nextdoor.com", "thumbtack.com", "google.com",
    "g.page", "goo.gl", "maps.app.goo.gl",
}

# Keyword hints only — the researcher confirms the real niche.
NICHE_KEYWORDS = {
    "house cleaning": ["cleaning", "maid", "housekeep", "janitorial", "deep clean", "cleaners",
                       "limpieza", "move out clean", "move-out clean"],
    "car detailing": ["detailing", "auto detail", "car detail", "mobile detail", "auto spa",
                      "ceramic coating", "paint correction", "car wash", "mobile wash"],
    "pest control": ["pest", "exterminat", "termite", "rodent", "bed bug", "bedbug",
                     "wildlife removal", "mosquito control"],
    "plumbing": ["plumb", "drain cleaning", "sewer", "water heater"],
    "hvac": ["hvac", "heating and cooling", "air conditioning", "furnace", "ac repair"],
    "tree service": ["tree service", "tree removal", "arborist", "stump", "tree trimming"],
    "water restoration": ["water damage", "mold remediation", "flood cleanup", "restoration"],
    "mobile mechanic": ["mobile mechanic", "mechanic", "auto repair"],
    "pet grooming": ["grooming", "groomer", "pet spa", "dog spa"],
    "roadside / towing": ["towing", "tow truck", "roadside", "jump start", "lockout"],
    "landscaping": ["landscap", "lawn care", "lawn service", "snow removal"],
    "handyman": ["handyman", "home repair"],
    "junk removal": ["junk removal", "junk hauling"],
    "roofing": ["roofing", "roofer"],
    "sunroom": ["sunroom", "patio enclosure"],
}

MARKETS = {
    "chicago, il": {
        "strong": ["chicago", "chicagoland", "chitown", "chi-town", "chi town", "evanston",
                   "oak park", "cicero", "skokie", "berwyn", "oak lawn", "evergreen park",
                   "des plaines", "park ridge", "niles", "lincolnwood", "harwood heights",
                   "norridge", "blue island", "calumet city", "burbank", "naperville",
                   "schaumburg", "aurora", "joliet", "elgin", "orland park", "tinley park",
                   "arlington heights", "palatine", "bolingbrook", "downers grove", "wheaton",
                   "lombard", "glenview", "wilmette", "elmhurst", "hinsdale", "la grange"],
        "weak": [r"\bil\b", r"\billinois\b"],
        "area_codes": ["312", "773", "872", "708", "847", "224", "630", "331"],
    },
}

EMAIL_RE = re.compile(r"[A-Za-z0-9._%+'-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
URL_RE = re.compile(
    r"(https?://[^\s,;\"'<>]+|www\.[^\s,;\"'<>]+"
    r"|\b(?:[a-z0-9-]+\.)+(?:com|net|org|co|us|biz|info|site|io|pro|me|app|services|cleaning|care"
    r"|online|store|shop|llc|business|company)\b(?:/[^\s,;\"'<>]*)?"
    # Link pages and booking sites often use other TLDs (linktr.ee, beacons.ai, wa.me).
    r"|\b(?:" + "|".join(re.escape(d) for d in sorted(LINK_PAGE_DOMAINS | BOOKING_DOMAINS)) + r")(?:/[^\s,;\"'<>]*)?)",
    re.I,
)
PHONE_RE = re.compile(r"(?:\+?1[\s.-]?)?\(?(\d{3})\)?[\s.-]?(\d{3})[\s.-]?(\d{4})")


def _key(header):
    return re.sub(r"[^a-z0-9]+", "_", str(header).strip().lower()).strip("_")


ALIAS_LOOKUP = {alias: field for field, aliases in FIELD_ALIASES.items() for alias in aliases}


def read_rows(path):
    """Return the export as a list of {header: value} dicts."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        sheet = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
        rows = [r for r in rows if any(c not in (None, "") for c in r)]
        if not rows:
            return []
        headers = [str(h or f"column_{i}") for i, h in enumerate(rows[0])]
        return [{h: ("" if v is None else v) for h, v in zip(headers, r)} for r in rows[1:]]
    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        if isinstance(data, dict):
            data = next((v for v in data.values() if isinstance(v, list)), [data])
        return [r for r in data if isinstance(r, dict)]
    if suffix == ".jsonl":
        lines = path.read_text(encoding="utf-8-sig").splitlines()
        return [json.loads(line) for line in lines if line.strip()]
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel_tab if suffix == ".tsv" else csv.excel
    return list(csv.DictReader(text.splitlines(), dialect=dialect))


def _clean(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _number(value):
    digits = re.sub(r"[^\d.]", "", str(value or ""))
    multiplier = 1
    text = str(value or "").strip().lower()
    if text.endswith("k"):
        multiplier = 1_000
    elif text.endswith("m"):
        multiplier = 1_000_000
    try:
        return int(float(digits) * multiplier) if digits else None
    except ValueError:
        return None


def domain_of(url):
    if not url:
        return ""
    if not re.match(r"https?://", url, re.I):
        url = "https://" + url
    host = urlparse(url).netloc.lower().split(":")[0]
    return host[4:] if host.startswith("www.") else host


def _matches(domain, domains):
    return any(domain == d or domain.endswith("." + d) for d in domains)


def website_kind(url):
    domain = domain_of(url)
    if not domain:
        return "none"
    if _matches(domain, LINK_PAGE_DOMAINS):
        return "link_page"
    if _matches(domain, BOOKING_DOMAINS):
        return "booking_page"
    if _matches(domain, SOCIAL_DOMAINS):
        return "social"
    return "own_site"


def classify_email(email):
    if not email:
        return "none"
    domain = email.split("@")[-1].lower()
    return "free_mail" if domain in FREE_MAIL_DOMAINS else "business_domain"


def format_phone(text):
    match = PHONE_RE.search(text or "")
    if not match:
        return "", ""
    area, prefix, line = match.groups()
    return f"({area}) {prefix}-{line}", area


def detect_niches(*texts):
    blob = " ".join(t.lower() for t in texts if t)
    hits = {}
    for niche, words in NICHE_KEYWORDS.items():
        count = sum(blob.count(w) for w in words)
        if count:
            hits[niche] = count
    return [n for n, _ in sorted(hits.items(), key=lambda kv: -kv[1])]


def market_config(market):
    key = market.strip().lower()
    if key in MARKETS:
        return MARKETS[key]
    city, _, state = key.partition(",")
    state = state.strip()
    return {
        "strong": [city.strip()],
        "weak": [rf"\b{re.escape(state)}\b"] if state else [],
        "area_codes": [],
    }


def location_hint(market, area_code, *texts):
    config = market_config(market)
    blob = " ".join(t.lower() for t in texts if t)
    if any(word in blob for word in config["strong"]):
        return "strong"
    if any(re.search(pattern, blob) for pattern in config["weak"]):
        return "weak"
    if area_code and area_code in config["area_codes"]:
        return "weak"
    return "none"


def map_row(row):
    mapped, extra = {}, {}
    for header, value in row.items():
        field = ALIAS_LOOKUP.get(_key(header))
        value = _clean(value)
        if not field:
            if value:
                extra[str(header)] = value
            continue
        if field == "location":
            if value:
                mapped["location"] = ", ".join(filter(None, [mapped.get("location"), value]))
        elif value and not mapped.get(field):
            mapped[field] = value
    # A bare "url" column is the Instagram profile only if it points at Instagram.
    profile = mapped.get("profile_url", "")
    if profile and "instagram.com" not in profile.lower():
        mapped.setdefault("website", profile)
        mapped["profile_url"] = ""
    return mapped, extra


def build_lead(index, row, market):
    mapped, extra = map_row(row)
    handle = mapped.get("handle", "").lstrip("@").strip().lower()
    profile_url = mapped.get("profile_url", "")
    if not handle and "instagram.com/" in profile_url.lower():
        handle = profile_url.rstrip("/").split("/")[-1].lstrip("@").lower()
    if handle and not profile_url:
        profile_url = f"https://www.instagram.com/{handle}/"
    bio = mapped.get("bio", "")

    emails = EMAIL_RE.findall(mapped.get("email", "")) or EMAIL_RE.findall(bio)
    email = emails[0].lower().rstrip(".") if emails else ""

    website = mapped.get("website", "")
    if not website:
        bio_urls = URL_RE.findall(EMAIL_RE.sub(" ", bio))
        website = bio_urls[0] if bio_urls else ""
    if website and not re.match(r"https?://", website, re.I):
        website = "https://" + website

    phone, area_code = format_phone(mapped.get("phone", "")) if mapped.get("phone") else format_phone(bio)

    niches = detect_niches(mapped.get("name", ""), bio, mapped.get("category", ""), handle)
    hint = location_hint(market, area_code, bio, mapped.get("location", ""), mapped.get("name", ""))
    kind = website_kind(website)
    email_type = classify_email(email)

    priority = 0
    priority += {"strong": 3, "weak": 1, "none": 0}[hint]
    priority += 2 if niches else 0
    priority += 1 if email_type == "business_domain" else 0
    priority += 1 if kind in ("own_site", "booking_page") else 0
    priority += 1 if phone else 0
    has_identity = any([mapped.get("name"), bio, website, email, phone])
    priority -= 0 if has_identity else 3

    return {
        "lead_id": f"L{index:04d}",
        "handle": handle,
        "profile_url": profile_url,
        "name": mapped.get("name", ""),
        "bio": bio,
        "category": mapped.get("category", ""),
        "location_text": mapped.get("location", ""),
        "email": email,
        "email_type": email_type,
        "email_domain": email.split("@")[-1] if email else "",
        "website": website,
        "website_kind": kind,
        "website_domain": domain_of(website) if kind == "own_site" else "",
        "phone": phone,
        "area_code": area_code,
        "followers": _number(mapped.get("followers")),
        "posts": _number(mapped.get("posts")),
        "is_business": mapped.get("is_business", ""),
        "niche_guess": niches[0] if niches else "",
        "niche_candidates": niches,
        "location_hint": hint,
        "priority": priority,
        "screen": "research",
        "duplicate_of": "",
        "extra": extra,
    }


def normalize(rows, market):
    """Build lead records, mark duplicates, and return them in file order."""
    leads, seen = [], {}
    for index, row in enumerate(rows, start=1):
        lead = build_lead(index, row, market)
        if not any([lead["handle"], lead["name"], lead["email"], lead["website"], lead["phone"]]):
            continue  # blank row
        keys = [
            ("handle", lead["handle"]),
            ("email", lead["email"]),
            ("website", lead["website_domain"]),
            ("phone", re.sub(r"\D", "", lead["phone"])),
        ]
        duplicate = next((seen[k] for k in keys if k[1] and k in seen), "")
        if duplicate:
            lead["screen"] = "skip_duplicate"
            lead["duplicate_of"] = duplicate
        else:
            for k in keys:
                if k[1]:
                    seen[k] = lead["lead_id"]
        leads.append(lead)
    return leads

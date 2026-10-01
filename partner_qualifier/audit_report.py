"""Partner audit outputs: a workbook for sorting and filtering, and an audit
book (Markdown) with the full write-up and email draft for each partner."""
from datetime import date

from openpyxl import Workbook

from .audit import local_leaders, money_keywords
from .report import _sheet

QUADRANT_FILL = {
    "Ideal partner: great work, under-marketed": "C6EFCE",
    "Good partner": "E2EFDA",
    "Strong and visible: may not need our jobs": "FFF2CC",
    "Vet first": "F8CBAD",
}


def _google(row):
    g = row["google"]
    return f"{g['rating']}★ ({g['reviews'] or '?'})" if g else "not confirmed"


def _first(points):
    return points[0]["point"] if points else ""


def build_audit_workbook(path, rows, markets):
    wb = Workbook()
    wb.remove(wb.active)

    best = []
    for rank, r in enumerate(rows, start=1):
        seo = r["audit"].get("seo") or {}
        best.append([
            rank, r["business_name"], r["quadrant"], r["quality"], r["need"],
            f"{r['visibility']}{'*' if r['visibility_estimated'] else ''}", _google(r), r["reviews_total"],
            r["website"] or "none", seo.get("est_monthly_traffic"),
            sum(1 for k in (r["audit"].get("keywords") or {}).get("ranking") or [] if (k.get("position") or 99) <= 10),
            ", ".join(h["keyword"] for h in r["pack_hits"]) or "no",
            _first(r["doing_right"]), _first(r["doing_wrong"]), r["email"]["status"], r["contact"].get("best", ""),
        ])
    ws = _sheet(wb, "Best Partners",
                ["Rank", "Business", "Fit", "Quality /10", "Need /10", "Visibility /10", "Google", "Reviews",
                 "Website", "Visitors/mo (est.)", "Page-1 money keywords", "In Google map pack for",
                 "Biggest strength", "Biggest gap", "Email", "Best contact"],
                best, widths={"Business": 28, "Fit": 30, "Biggest strength": 45, "Biggest gap": 45,
                              "Best contact": 34, "In Google map pack for": 30, "Website": 28})
    from openpyxl.styles import PatternFill
    for row in ws.iter_rows(min_row=2):
        fill = QUADRANT_FILL.get(row[2].value)
        if fill:
            row[2].fill = PatternFill("solid", fgColor=fill)

    seo_rows = []
    for r in rows:
        seo = r["audit"].get("seo") or {}
        authority = seo.get("authority") or {}
        issues = seo.get("technical_issues") or []
        ranking = sorted((r["audit"].get("keywords") or {}).get("ranking") or [], key=lambda k: k.get("position") or 99)
        seo_rows.append([
            r["business_name"], r["website"] or "none", seo.get("organic_keywords"), seo.get("top10_keywords"),
            seo.get("top3_keywords"), seo.get("est_monthly_traffic"), seo.get("traffic_value_usd"),
            authority.get("domain_rank"), authority.get("referring_domains"), seo.get("onpage_score"),
            "; ".join(f"{i.get('issue')} ({i.get('severity')})" for i in issues[:3]),
            "; ".join(f"{k.get('keyword')} #{k.get('position')}" for k in ranking[:8]),
            "; ".join(f"{k['keyword']} ({k['volume'] or '?'}/mo)" for k in r["missing_keywords"][:5]),
        ])
    _sheet(wb, "SEO & Keywords",
           ["Business", "Website", "Ranking keywords", "Top-10 keywords", "Top-3 keywords", "Visitors/mo (est.)",
            "Traffic value $/mo", "Authority (0-1000)", "Linking sites", "On-page score", "Top technical issues",
            "Where they rank (money keywords)", "Biggest searches they miss"],
           seo_rows, widths={"Business": 28, "Where they rank (money keywords)": 50,
                             "Biggest searches they miss": 50, "Top technical issues": 36})

    leader_rows, demand_rows = [], []
    for niche, market in markets.items():
        for b in local_leaders(market):
            leader_rows.append([niche, b["name"], b["domain"], b["rating"], b["reviews"], b["appearances"],
                                ", ".join(b["keywords"])])
        for k in money_keywords(market):
            demand_rows.append([niche, k.get("keyword"), k.get("volume"), k.get("cpc"), k.get("intent")])
    _sheet(wb, "Market Leaders", ["Niche", "Business", "Website", "Google rating", "Reviews",
                                  "Map-pack appearances", "Searches"], leader_rows,
           widths={"Business": 30, "Searches": 60})
    _sheet(wb, "Demand Keywords", ["Niche", "Search", "Monthly searches", "Ad cost per click $", "Intent"],
           demand_rows, widths={"Search": 40})

    rw_rows = []
    for r in rows:
        for kind, points in (("Doing right", r["doing_right"]), ("Doing wrong", r["doing_wrong"])):
            for p in points:
                rw_rows.append([r["business_name"], kind, p.get("point"), p.get("evidence"), p.get("source")])
    _sheet(wb, "Right & Wrong", ["Business", "Type", "Point", "Evidence", "Source"], rw_rows,
           widths={"Business": 28, "Point": 50, "Evidence": 60, "Source": 50})

    email_rows = []
    for rank, r in enumerate(rows, start=1):
        e = r["email"]
        email_rows.append([rank, r["business_name"], r["contact"].get("email") or r["contact"].get("best", ""),
                           e.get("subject", ""), e["full_text"], e["status"], "; ".join(e["failed"] + e["warnings"]),
                           e["grade"]])
    _sheet(wb, "Email Drafts", ["Rank", "Business", "To", "Subject", "Email", "Status", "Problems",
                                "Reading grade"], email_rows,
           widths={"Business": 28, "To": 32, "Subject": 34, "Email": 80, "Problems": 40})
    wb.save(path)


def _src(item):
    source = (item or {}).get("source") or ""
    return f" ([source]({source}))" if source.startswith("http") else f" ({source})" if source else ""


def build_audit_book(path, rows, markets, run):
    out = [f"# Partner Audit Book: {run['market']}",
           f"*{len(rows)} partners audited · {date.today().isoformat()} · run `{run['run_id']}`*", "",
           "**How to read this:** **Quality** (0–10) means we can trust them with our customers. "
           "**Need** (0–10) means our jobs would matter to them; higher means they're harder to find "
           "online today. **Visibility** is the opposite of need, and `*` means part of it was estimated. "
           "Best partners: high quality, high need.", "",
           "| # | Business | Fit | Quality | Need | Google | Email |", "|---|---|---|---|---|---|---|"]
    for rank, r in enumerate(rows, start=1):
        out.append(f"| {rank} | [{r['business_name']}](#{rank}) | {r['quadrant']} | {r['quality']}/10 | "
                   f"{r['need']}/10 | {_google(r)} | {r['email']['status']} |")

    for niche, market in markets.items():
        leaders = local_leaders(market)[:6]
        if leaders:
            out += ["", f"## Market: {niche} in {run['market']}", "",
                    "Who wins Google's map pack (top 3) for the money searches:", "",
                    "| Business | Google | Reviews | Map-pack appearances |", "|---|---|---|---|"]
            out += [f"| {b['name']} | {b['rating']}★ | {b['reviews'] or '?'} | {b['appearances']} |" for b in leaders]
            top = money_keywords(market)[:6]
            if top:
                out += ["", "Biggest money searches: " + "; ".join(f"{k['keyword']} ({k.get('volume') or '?'}/mo)" for k in top)]

    for rank, r in enumerate(rows, start=1):
        a, seo = r["audit"], r["audit"].get("seo") or {}
        rep = a.get("google_reputation") or {}
        authority = seo.get("authority") or {}
        out += ["", "---", "", f'<a id="{rank}"></a>', f"## {rank}. {r['business_name']}",
                f"**{r['quadrant']}** · Quality {r['quality']}/10 · Need {r['need']}/10 · "
                f"Visibility {r['visibility']}/10{'*' if r['visibility_estimated'] else ''} · {r['niche'] or ''}", "",
                f"Contact: {r['contact'].get('best', '')} · Website: {r['website'] or 'none'}"]

        out += ["", "### Google reputation"]
        g = r["google"]
        out.append(f"- **Google:** {g['rating']}★ from {g['reviews'] or '?'} reviews, via {g['via']}{_src(g)}"
                   if g else "- **Google:** rating not confirmed")
        for p in rep.get("other_platforms") or []:
            out.append(f"- **{p.get('platform')}:** {p.get('rating')}★ ({p.get('count')}){_src(p)}")
        for label, key in (("Customers praise", "praise_themes"), ("Customers complain about", "complaint_themes")):
            themes = rep.get(key) or []
            if themes:
                out.append(f"- **{label}:** " + "; ".join(f"{t.get('theme')}{_src(t)}" for t in themes))
        if rep.get("owner_replies_to_reviews") is not None:
            out.append(f"- **Replies to reviews:** {'yes' if rep['owner_replies_to_reviews'] else 'no'}"
                       f"{_src({'source': rep.get('replies_source')})}")
        out.append("- **Google map pack:** " + (", ".join(f"in top 3 for \"{h['keyword']}\"" for h in r["pack_hits"])
                                                 or "not in the top 3 for any money search checked"))

        out += ["", "### SEO"]
        if seo.get("website_type") in ("none", "social_only", "link_page"):
            out.append(f"- No real website ({seo.get('website_type')}). Customers can't find them in Google search "
                       "results, only through Instagram, directories or word of mouth.")
        elif seo.get("organic_keywords") is not None:
            out += [f"- Ranks for **{seo.get('organic_keywords'):,}** searches: {seo.get('top10_keywords')} on page 1, "
                    f"{seo.get('top3_keywords')} in the top 3",
                    f"- About **{(seo.get('est_monthly_traffic') or 0):,} visitors/month** (estimate), "
                    f"worth about ${(seo.get('traffic_value_usd') or 0):,.0f}/month in ads"]
            if authority:
                out.append(f"- Authority {authority.get('domain_rank')}/1000 · {authority.get('referring_domains')} "
                           f"linking sites · spam score {authority.get('spam_score')}")
            if seo.get("onpage_score") is not None:
                issues = seo.get("technical_issues") or []
                out.append(f"- Site health {seo['onpage_score']}/100" +
                           (": " + "; ".join(f"{i.get('issue')} ({i.get('severity')})" for i in issues[:3]) if issues else ""))
            for page in (seo.get("top_pages") or [])[:3]:
                out.append(f"- Top page: {page.get('url')}, ~{page.get('traffic')}/mo"
                           + (f": {page['note']}" if page.get("note") else ""))
            out.append(f"- Source: {seo.get('source', 'OpenRush')}")
        else:
            out.append("- Not checked")

        out += ["", "### Keywords"]
        ranking = sorted((a.get("keywords") or {}).get("ranking") or [], key=lambda k: k.get("position") or 99)
        if ranking:
            out += ["| Search | Position | Monthly searches |", "|---|---|---|"]
            out += [f"| {k.get('keyword')} | #{k.get('position')} | {k.get('volume') or '?'} |" for k in ranking[:10]]
        else:
            out.append("- No rankings found for this market's money searches")
        if r["missing_keywords"]:
            out.append("- **Biggest searches they're missing:** " +
                       "; ".join(f"{k['keyword']} ({k['volume'] or '?'}/mo)" for k in r["missing_keywords"][:6]))

        comp = a.get("competition") or {}
        out += ["", "### Competition"]
        if comp.get("position_summary"):
            out.append(f"- {comp['position_summary']}")
        for c in (comp.get("organic_competitors") or [])[:5]:
            out.append(f"- Competes with **{c.get('domain')}**" +
                       (f" ({c.get('shared_keywords')} shared searches)" if c.get("shared_keywords") else ""))

        for title, points in (("What they're doing right", r["doing_right"]), ("What they're doing wrong", r["doing_wrong"])):
            out += ["", f"### {title}"]
            out += [f"- **{p.get('point')}**: {p.get('evidence', '')}{_src(p)}" for p in points] or ["- None sourced"]

        signals = a.get("capacity_signals") or []
        if signals:
            out += ["", "### Signs they want more work"]
            out += [f"- {s.get('signal')}{_src(s)}" for s in signals]

        angles = a.get("email_angles") or []
        if angles:
            out += ["", "### Other email angles"]
            out += [f"- **{x.get('angle')}:** {x.get('hook')}{_src(x)}" for x in angles]

        e = r["email"]
        out += ["", f"### Custom email: {e['status']}", f"**To:** {r['contact'].get('email') or r['contact'].get('best', '')}  ",
                f"**Subject:** {e.get('subject', '')}", "", "```", e["full_text"], "```"]
        if e["failed"] or e["warnings"]:
            out.append("Fix before sending: " + "; ".join(e["failed"] + e["warnings"]))
        facts = e.get("facts_used") or []
        if facts:
            out.append("Facts used: " + "; ".join(f"{f.get('fact')}{_src(f)}" for f in facts))
    path.write_text("\n".join(out) + "\n")
